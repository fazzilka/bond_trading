import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import xlwt
from numbers_parser import Document

from bond_trading.infrastructure.imports import XlsxImportError, XlsxPortfolioReader

HEADERS = [
    None,
    "ISIN",
    "Наименование",
    "Оферта",
    "Дата приобретения",
    "дата подачи на оферту",
    "дата оферты",
    "Кол-во",
    "цена приобретения",
    "НКД",
    "нкд сумма",
    "комиссия",
    "стоимость приобретения",
    "цена к погашению",
    "купон",
    "Кол-во купонов",
    "сумма купонов",
    "комиссия",
    "стоимость к погашению",
    "доход до налога",
    "чистый доход",
    "Годовой %, план",
    "Годовой текущий %",
]


def workbook_bytes(rows: list[list[object]], *, headers: bool = True) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Доход счёт 2026"
    worksheet.append([None])
    if headers:
        worksheet.append(HEADERS)
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def lot_row(isin: str, quantity: object = 40, purchase_date: object | None = None) -> list[object]:
    row: list[object] = [None] * 23
    row[1] = isin
    row[2] = "ЭкономЛизинг 1р-07"
    row[3] = "погашение"
    row[4] = purchase_date or datetime(2026, 5, 25)
    row[6] = datetime(2027, 2, 15)
    row[7] = quantity
    row[8] = 962.9
    row[9] = 3.51
    row[11] = 0.39
    row[13] = 1000
    row[21] = 19.3
    return row


def test_reads_main_block_and_ignores_control_block() -> None:
    content = workbook_bytes(
        [
            lot_row("RU000A107SX3", 40),
            lot_row(" RU000A107sg8 ", 8),
            lot_row("RU000A107SX3", 2),
            [None] * 23,
            lot_row("RU000A107SX3", 40),
        ]
    )

    preview = XlsxPortfolioReader().preview(content, "portfolio.xlsx")

    assert preview.header_row_number == 2
    assert preview.rows_read == 3
    assert len(preview.rows) == 3
    assert preview.rows[1].normalized_isin == "RU000A107SG8"
    assert preview.rows[0].normalized_isin == preview.rows[2].normalized_isin
    assert preview.rows[0].quantity == Decimal("40")
    assert preview.rows[0].purchase_commission_rub_per_bond == Decimal("0.39")
    assert preview.errors == ()


def test_partial_errors_do_not_drop_valid_rows() -> None:
    content = workbook_bytes(
        [
            lot_row("RU000A107SX3"),
            lot_row("RU000A107SX3", 0),
            lot_row("RU000A107SX3", purchase_date="not-a-date"),
            lot_row("RU000A107SX4"),
        ]
    )

    preview = XlsxPortfolioReader().preview(content, "portfolio.xlsx")

    assert len(preview.rows) == 1
    assert [error.field for error in preview.errors] == ["quantity", "purchase_date", "isin"]
    assert [error.row_number for error in preview.errors] == [4, 5, 6]


def test_manual_formula_is_rejected() -> None:
    content = workbook_bytes([lot_row("RU000A107SX3", quantity="=20+20")])

    preview = XlsxPortfolioReader().preview(content, "portfolio.xlsx")

    assert preview.rows == ()
    assert preview.errors[0].field == "quantity"
    assert "formula" in preview.errors[0].message


def test_missing_header_and_wrong_extension() -> None:
    reader = XlsxPortfolioReader()
    with pytest.raises(XlsxImportError, match="ISIN"):
        reader.preview(workbook_bytes([], headers=False), "portfolio.xlsx")
    with pytest.raises(XlsxImportError, match="extensions"):
        reader.preview(b"not a workbook", "portfolio.csv")


def test_iso_date_and_decimal_comma_are_supported() -> None:
    row = lot_row("RU000A107SX3", quantity="40,5", purchase_date="2026-05-25")
    preview = XlsxPortfolioReader().preview(workbook_bytes([row]), "portfolio.xlsx")

    assert preview.rows[0].purchase_date == date(2026, 5, 25)
    assert preview.rows[0].quantity == Decimal("40.5")


def test_legacy_xls_is_supported() -> None:
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Доход счёт 2026")
    for column, value in enumerate(HEADERS):
        if value is not None:
            worksheet.write(1, column, value)
    for column, value in enumerate(lot_row("RU000A107SX3")):
        if value is not None:
            if isinstance(value, datetime):
                worksheet.write(
                    2,
                    column,
                    value,
                    xlwt.easyxf(num_format_str="YYYY-MM-DD"),
                )
            else:
                worksheet.write(2, column, value)
    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = XlsxPortfolioReader().preview(buffer.getvalue(), "portfolio.xls")

    assert preview.sheet_name == "Доход счёт 2026"
    assert len(preview.rows) == 1
    assert preview.rows[0].normalized_isin == "RU000A107SX3"


def test_apple_numbers_is_supported(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.numbers"
    document = Document()
    document.add_sheet("Доход счёт 2026", "Таблица 1")
    table = document.sheets["Доход счёт 2026"].tables[0]
    for column, value in enumerate(HEADERS):
        if value is not None:
            table.write(1, column, value)
    for column, value in enumerate(lot_row("RU000A107SX3")):
        if value is not None:
            table.write(2, column, value)
    for column, value in enumerate(lot_row("RU000A107SX3", quantity=999)):
        if value is not None:
            table.write(4, column, value)
    document.save(path)

    preview = XlsxPortfolioReader().preview(path.read_bytes(), path.name)

    assert preview.sheet_name == "Доход счёт 2026"
    assert len(preview.rows) == 1
    assert preview.rows[0].quantity == Decimal("40")
